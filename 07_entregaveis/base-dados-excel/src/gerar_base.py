"""Gera a base de dados do benchmark (BIA-ACL-BD-01) a partir do repositório.

Entradas
- 02_frente-1_benchmark-estrategico/fichas-benchmark/*.md  cabeçalho YAML e seções de cada ficha
- 06_fontes/catalogo-fontes.csv                           catálogo de fontes
- 03_frente-2_maturidade-roadmaps/regua-classificacao.md   tabelas marcadas com <!-- tabela:… -->

Saída
- 07_entregaveis/base-dados-excel/base-dados-benchmark.xlsx

Uso, na raiz do repositório:
    uv run --with openpyxl --with pyyaml python 07_entregaveis/base-dados-excel/src/gerar_base.py
"""

import csv
import datetime as dt
import re
import sys
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parents[3]
FICHAS = RAIZ / "02_frente-1_benchmark-estrategico" / "fichas-benchmark"
CATALOGO = RAIZ / "06_fontes" / "catalogo-fontes.csv"
REGUA = RAIZ / "03_frente-2_maturidade-roadmaps" / "regua-classificacao.md"
SAIDA = RAIZ / "07_entregaveis" / "base-dados-excel" / "base-dados-benchmark.xlsx"

LEGENDA = {
    "Contratante": "Acelen Energia Renovável",
    "Documento": "Base de dados do benchmark — Benchmarking de uso de IA",
    "Nº / Revisão": "BIA-ACL-BD-01 · Rev. 00",
    "Responsáveis": "Dr. Guilherme M. Sanches · Rafael Rotiroti — OagronomIA",
}
RODAPE = "Acelen Energia Renovável · Base de dados do benchmark · BIA-ACL-BD-01 · Rev. 00"

ESTAGIOS = ["Exploração", "Pilotos", "Operacionalização", "Escala", "Transformação"]
DIMENSOES = [("governanca", "Governança"), ("dados", "Dados"), ("organizacao", "Organização"), ("aplicacoes", "Aplicações")]

# Design system OagronomIA, versão clara para exportação
VERDE = "52C937"
VERDE_ESCURO = "2F8B1F"
FUNDO_CABECALHO = "0B100B"
PAPEL = "F4F1EA"
LINHA = "DCD7C7"
CINZA = "7C8678"
CORPO = "Manrope"
DISPLAY = "Chakra Petch"
MONO = "JetBrains Mono"

# ------------------------------------------------------------------ leitura


def ler_fichas():
    fichas = []
    for arquivo in sorted(FICHAS.glob("*.md")):
        if arquivo.name.startswith("_"):
            continue
        texto = arquivo.read_text(encoding="utf-8")
        partes = texto.split("---", 2)
        if not texto.startswith("---") or len(partes) < 3:
            sys.exit(f"{arquivo.name}: ficha sem cabeçalho YAML entre linhas ---")
        cabecalho = yaml.safe_load(partes[1]) or {}
        corpo = partes[2]
        cabecalho["_arquivo"] = arquivo.relative_to(RAIZ).as_posix()
        cabecalho["_sintese"] = extrair_sintese(corpo)
        cabecalho["_funcionou"] = extrair_itens(corpo, "O que funcionou")
        cabecalho["_nao_funcionou"] = extrair_itens(corpo, "O que não funcionou ou travou")
        cabecalho["_licoes"] = extrair_itens(corpo, "Lições para a leitura transversal")
        cabecalho["_lacunas"] = extrair_itens(corpo, "Lacunas de evidência")
        validar_ficha(cabecalho)
        fichas.append(cabecalho)
    return fichas


def extrair_sintese(corpo):
    for linha in corpo.splitlines():
        if linha.startswith(">"):
            return limpar(linha.lstrip("> "))
    return ""


def extrair_itens(corpo, titulo):
    itens, dentro = [], False
    for linha in corpo.splitlines():
        if linha.startswith("#"):
            dentro = linha.lstrip("# ").strip() == titulo
            continue
        if dentro and linha.startswith("- "):
            item = limpar(linha[2:])
            if item and item != "…":
                itens.append(item)
    return "\n".join(itens)


def limpar(texto):
    return texto.replace("**", "").strip()


def validar_ficha(f):
    erros = []
    for chave, _ in DIMENSOES:
        dim = (f.get("notas") or {}).get(chave) or {}
        nota, conf = dim.get("nota"), dim.get("confianca")
        if nota is not None and nota not in (1, 2, 3, 4, 5):
            erros.append(f"nota de {chave} fora de 1–5: {nota}")
        if nota is not None and conf not in ("A", "B", "C"):
            erros.append(f"nota de {chave} sem confiança A, B ou C")
        if nota in (4, 5) and conf != "A":
            erros.append(f"nota {nota} em {chave} exige confiança A")
    for lente, nivel in (f.get("lentes") or {}).items():
        if nivel is not None and nivel not in (0, 1, 2, 3, 4, 5):
            erros.append(f"lente {lente} fora de 0–5: {nivel}")
    if erros:
        sys.exit(f"{f['_arquivo']}: " + "; ".join(erros))


def ler_catalogo():
    with CATALOGO.open(encoding="utf-8", newline="") as arq:
        return list(csv.DictReader(arq))


def ler_tabela(marcador):
    """Devolve as linhas da tabela markdown logo após <!-- tabela:marcador -->."""
    linhas = REGUA.read_text(encoding="utf-8").splitlines()
    alvo = f"<!-- tabela:{marcador} -->"
    if alvo not in linhas:
        sys.exit(f"regua-classificacao.md: marcador {alvo} não encontrado")
    tabela = []
    for linha in linhas[linhas.index(alvo) + 1:]:
        if not linha.startswith("|"):
            if tabela:
                break
            continue
        celulas = [c.strip() for c in linha.strip().strip("|").split("|")]
        if all(set(c) <= set("-: ") for c in celulas):
            continue
        tabela.append(celulas)
    return tabela


# ------------------------------------------------------------------ estilo

fonte_corpo = Font(name=CORPO, size=10)
fonte_cabecalho = Font(name=CORPO, size=10, bold=True, color="FFFFFF")
fonte_titulo = Font(name=DISPLAY, size=18, bold=True, italic=True, color=VERDE_ESCURO)
fonte_rotulo = Font(name=DISPLAY, size=9, bold=True, color=CINZA)
fonte_secao = Font(name=DISPLAY, size=12, bold=True, color=VERDE_ESCURO)
fonte_mono = Font(name=MONO, size=10)
preenchimento_cabecalho = PatternFill("solid", fgColor=FUNDO_CABECALHO)
preenchimento_exemplo = PatternFill("solid", fgColor=PAPEL)
borda = Border(bottom=Side(style="thin", color=LINHA))
quebra = Alignment(wrap_text=True, vertical="top")


def preparar_aba(ws, cabecalho, larguras):
    ws.append(cabecalho)
    for col, largura in enumerate(larguras, start=1):
        celula = ws.cell(row=1, column=col)
        celula.font = fonte_cabecalho
        celula.fill = preenchimento_cabecalho
        celula.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = largura
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    ws.oddFooter.center.text = RODAPE
    ws.oddFooter.center.size = 8
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def formatar_corpo(ws, colunas_mono=()):
    for linha in ws.iter_rows(min_row=2):
        for celula in linha:
            celula.font = fonte_mono if celula.column in colunas_mono else fonte_corpo
            celula.alignment = quebra
            celula.border = borda
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions


# ------------------------------------------------------------------ abas


def aba_leia_me(wb, fichas, fontes, gerado_em):
    ws = wb.active
    ws.title = "Leia-me"
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 100
    ws["A1"] = "Base de dados do benchmark"
    ws["A1"].font = fonte_titulo
    ws.row_dimensions[1].height = 30
    linha = 3
    for rotulo, valor in {**LEGENDA, "Gerado em": gerado_em}.items():
        ws.cell(row=linha, column=1, value=rotulo.upper()).font = fonte_rotulo
        ws.cell(row=linha, column=2, value=valor).font = fonte_corpo
        linha += 1

    blocos = [
        ("Abas", [
            ("Painel", "Empresas do painel: identificação, setor, porte e situação da ficha."),
            ("Classificação", "Notas de 1 a 5 por dimensão, com confiança A, B ou C, e o estágio consolidado calculado por fórmula. Lentes de IA generativa e campo de 0 a 5."),
            ("Fichas", "Síntese, casos, lições e lacunas de cada ficha, com os IDs das fontes."),
            ("Régua", "Estágios e critérios de cada nível, por dimensão e por lente."),
            ("Fontes", "Catálogo de fontes. Toda afirmação do estudo aponta para um ID desta aba."),
        ]),
        ("Regras", [
            ("Sem evidência, sem nota", "Célula de nota vazia significa falta de evidência pública, não nota 1."),
            ("Confiança", "A: fonte primária (a própria empresa ou o autor do estudo). B: secundária qualificada. C: indireta. Notas 4 e 5 exigem A."),
            ("Estágio consolidado", "Mediana das quatro notas, limitada pela nota de Dados + 1. Exige ao menos três notas, entre elas a de Dados. Mediana entre dois estágios indica transição (ex.: 3→4, valor 3,5)."),
            ("Níveis das lentes", "0 = sem evidência; 1 a 5 conforme a aba Régua. As lentes não entram no estágio consolidado."),
        ]),
        ("Como atualizar", [
            ("Não edite à mão", "Esta base é gerada a partir das fichas (02_frente-1_benchmark-estrategico/fichas-benchmark), do catálogo (06_fontes/catalogo-fontes.csv) e da régua (03_frente-2_maturidade-roadmaps/regua-classificacao.md)."),
            ("Comando", "uv run --with openpyxl --with pyyaml python 07_entregaveis/base-dados-excel/src/gerar_base.py"),
        ]),
    ]
    linha += 1
    for titulo, itens in blocos:
        ws.cell(row=linha, column=1, value=titulo).font = fonte_secao
        linha += 1
        for rotulo, texto in itens:
            ws.cell(row=linha, column=1, value=rotulo).font = Font(name=CORPO, size=10, bold=True)
            celula = ws.cell(row=linha, column=2, value=texto)
            celula.font = fonte_mono if rotulo == "Comando" else fonte_corpo
            celula.alignment = quebra
            linha += 1
        linha += 1

    resumo = f"{len(fichas)} ficha(s) · {len(fontes)} fonte(s) no catálogo"
    if not fichas:
        resumo += ". Ainda não há fichas: a aba Classificação traz uma linha de exemplo, marcada, só para mostrar o cálculo."
    ws.cell(row=linha, column=1, value="Situação").font = fonte_secao
    ws.cell(row=linha, column=2, value=resumo).font = fonte_corpo
    ws.oddFooter.center.text = RODAPE
    ws.oddFooter.center.size = 8


def aba_painel(wb, fichas):
    ws = wb.create_sheet("Painel")
    preparar_aba(ws, ["ID", "Empresa", "Grupo", "Setor", "Subsetor", "País", "Capital", "Porte",
                      "Situação da ficha", "Responsável", "Revisor", "Atualizado em", "Arquivo"],
                 [8, 28, 12, 16, 20, 12, 14, 36, 16, 12, 12, 14, 48])
    for f in fichas:
        ws.append([f.get("id"), f.get("empresa"), f.get("grupo"), f.get("setor"), f.get("subsetor"),
                   f.get("pais"), f.get("capital"), f.get("porte"), f.get("situacao"),
                   f.get("responsavel"), f.get("revisor"), str(f.get("atualizado_em") or ""), f["_arquivo"]])
    formatar_corpo(ws, colunas_mono=(1,))


def formula_estagio(r):
    """Fórmulas da linha r da aba Classificação (notas em D, F, H, J)."""
    notas = f"D{r},F{r},H{r},J{r}"
    nomes = ",".join(f'"{n}"' for n in ESTAGIOS)
    return {
        "L": f"=COUNT({notas})",
        "M": f'=IF(L{r}=0,"",MEDIAN({notas}))',
        "N": f'=IF(F{r}="","",F{r}+1)',
        "O": f'=IF(OR(L{r}<3,F{r}=""),"",MIN(M{r},N{r}))',
        "P": f'=IF(O{r}="","Evidência insuficiente",IF(O{r}=INT(O{r}),INT(O{r})&"",INT(O{r})&"→"&(INT(O{r})+1)))',
        "Q": f'=IF(O{r}="","",IF(O{r}=INT(O{r}),CHOOSE(O{r},{nomes}),CHOOSE(INT(O{r}),{nomes})&" → "&CHOOSE(INT(O{r})+1,{nomes})))',
    }


def aba_classificacao(wb, fichas):
    ws = wb.create_sheet("Classificação")
    preparar_aba(ws, ["ID", "Empresa", "Setor",
                      "Governança", "Conf.", "Dados", "Conf.", "Organização", "Conf.", "Aplicações", "Conf.",
                      "Notas", "Mediana", "Teto (Dados + 1)", "Estágio (valor)", "Estágio", "Nome do estágio",
                      "Lente IA generativa", "Lente campo"],
                 [8, 28, 16, 12, 7, 10, 7, 13, 7, 12, 7, 8, 10, 11, 11, 16, 34, 12, 12])
    linhas = []
    for f in fichas:
        notas = f.get("notas") or {}
        lentes = f.get("lentes") or {}
        valores = [f.get("id"), f.get("empresa"), f.get("setor")]
        for chave, _ in DIMENSOES:
            dim = notas.get(chave) or {}
            valores += [dim.get("nota"), dim.get("confianca")]
        linhas.append((valores, [lentes.get("ia_generativa"), lentes.get("campo")], False))
    if not linhas:
        exemplo = ["EX", "Exemplo ilustrativo — não faz parte do painel", "—", 3, "A", 3, "A", 4, "A", 4, "A"]
        linhas.append((exemplo, [3, 2], True))

    for i, (valores, lentes, eh_exemplo) in enumerate(linhas, start=2):
        for col, valor in enumerate(valores, start=1):
            ws.cell(row=i, column=col, value=valor)
        for col, formula in formula_estagio(i).items():
            ws[f"{col}{i}"] = formula
        ws.cell(row=i, column=18, value=lentes[0])
        ws.cell(row=i, column=19, value=lentes[1])
    formatar_corpo(ws, colunas_mono=(1, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16, 18, 19))
    for i, (_, _, eh_exemplo) in enumerate(linhas, start=2):
        if eh_exemplo:
            for col in range(1, 20):
                ws.cell(row=i, column=col).fill = preenchimento_exemplo
                ws.cell(row=i, column=col).font = Font(name=CORPO, size=10, italic=True, color=CINZA)

    ultima = ws.max_row
    escala = ColorScaleRule(start_type="num", start_value=1, start_color="EAE6DA",
                            mid_type="num", mid_value=3, mid_color="9BE05F",
                            end_type="num", end_value=5, end_color=VERDE)
    for col in ("D", "F", "H", "J", "O"):
        ws.conditional_formatting.add(f"{col}2:{col}{ultima}", escala)
    for col in ("M", "O"):
        for r in range(2, ultima + 1):
            ws[f"{col}{r}"].number_format = "0.0"


def aba_fichas(wb, fichas):
    ws = wb.create_sheet("Fichas")
    preparar_aba(ws, ["ID", "Empresa", "Síntese", "O que funcionou", "O que não funcionou ou travou",
                      "Lições", "Lacunas de evidência", "Fontes", "Arquivo"],
                 [8, 24, 50, 50, 50, 40, 36, 20, 44])
    for f in fichas:
        ws.append([f.get("id"), f.get("empresa"), f["_sintese"], f["_funcionou"], f["_nao_funcionou"],
                   f["_licoes"], f["_lacunas"], ", ".join(f.get("fontes") or []), f["_arquivo"]])
    formatar_corpo(ws, colunas_mono=(1, 8))


def aba_regua(wb):
    ws = wb.create_sheet("Régua")
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 60
    ws.column_dimensions["D"].width = 60
    ws.oddFooter.center.text = RODAPE
    ws.oddFooter.center.size = 8
    secoes = [
        ("Estágios", "estagios"),
        ("Governança de IA", "regua-governanca"),
        ("Dados", "regua-dados"),
        ("Organização", "regua-organizacao"),
        ("Aplicações", "regua-aplicacoes"),
        ("Lente: IA generativa", "regua-ia-generativa"),
        ("Lente: campo e sistemas autônomos", "regua-campo"),
    ]
    linha = 1
    for titulo, marcador in secoes:
        ws.cell(row=linha, column=1, value=titulo).font = fonte_secao
        linha += 1
        tabela = ler_tabela(marcador)
        cabecalho, corpo = tabela[0], tabela[1:]
        if len(cabecalho) == 3:  # Nível | Critério | Evidência: deixa a coluna B vazia para alinhar
            cabecalho = [cabecalho[0], "", cabecalho[1], cabecalho[2]]
            corpo = [[c[0], "", c[1], c[2]] for c in corpo]
        for col, valor in enumerate(cabecalho, start=1):
            celula = ws.cell(row=linha, column=col, value=valor)
            celula.font = fonte_cabecalho
            celula.fill = preenchimento_cabecalho
        linha += 1
        for registro in corpo:
            for col, valor in enumerate(registro, start=1):
                celula = ws.cell(row=linha, column=col, value=int(valor) if col == 1 and valor.isdigit() else valor)
                celula.font = fonte_mono if col == 1 else fonte_corpo
                celula.alignment = quebra
                celula.border = borda
            linha += 1
        linha += 1


def aba_fontes(wb, fontes):
    ws = wb.create_sheet("Fontes")
    colunas = [("id", "ID", 9), ("titulo", "Título", 48), ("autor_organizacao", "Autor/organização", 26),
               ("data_publicacao", "Data de publicação", 12), ("url", "URL", 44), ("tipo", "Tipo", 20),
               ("nivel", "Nível", 7), ("acesso", "Acesso", 14), ("empresas", "Empresas", 20),
               ("frentes", "Frentes", 9), ("usado_em", "Usado em", 24), ("acessado_em", "Acessado em", 12),
               ("observacoes", "Observações", 50)]
    preparar_aba(ws, [c[1] for c in colunas], [c[2] for c in colunas])
    for fonte in fontes:
        ws.append([fonte.get(c[0], "") for c in colunas])
    formatar_corpo(ws, colunas_mono=(1, 4, 7, 12))


def main():
    fichas = ler_fichas()
    fontes = ler_catalogo()
    ids_catalogo = {f["id"] for f in fontes}
    for f in fichas:
        faltando = [i for i in f.get("fontes") or [] if i not in ids_catalogo]
        if faltando:
            sys.exit(f"{f['_arquivo']}: fontes fora do catálogo: {', '.join(faltando)}")

    hoje = dt.date.today()
    wb = Workbook()
    aba_leia_me(wb, fichas, fontes, hoje.strftime("%d/%m/%Y"))
    aba_painel(wb, fichas)
    aba_classificacao(wb, fichas)
    aba_fichas(wb, fichas)
    aba_regua(wb)
    aba_fontes(wb, fontes)
    wb.calculation.fullCalcOnLoad = True
    momento = dt.datetime.combine(hoje, dt.time())
    wb.properties.creator = "OagronomIA"
    wb.properties.title = LEGENDA["Documento"]
    wb.properties.created = momento
    wb.properties.modified = momento
    SAIDA.parent.mkdir(parents=True, exist_ok=True)
    wb.save(SAIDA)
    print(f"{SAIDA.relative_to(RAIZ)}: {len(fichas)} ficha(s), {len(fontes)} fonte(s)")


if __name__ == "__main__":
    main()
